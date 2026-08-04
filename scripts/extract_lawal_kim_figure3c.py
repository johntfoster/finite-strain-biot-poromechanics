#!/usr/bin/env python3
"""Extract the exact Figure 3c source values from the authors' public workbook."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

import openpyxl


SERIES = {
    "SP14": {
        "label": "14-day serpentinized",
        "pressure_col": 6,
        "modulus_col": 7,
        "alpha_col": 9,
        "grain_modulus_cell": "G28",
    },
    "HT14": {
        "label": "14-day dry heat-treated control",
        "pressure_col": 14,
        "modulus_col": 15,
        "alpha_col": 17,
        "grain_modulus_cell": "O28",
    },
    "SP30": {
        "label": "30-day serpentinized",
        "pressure_col": 22,
        "modulus_col": 23,
        "alpha_col": 25,
        "grain_modulus_cell": "W28",
    },
}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    sheet = workbook["Integrated"]
    records: list[dict[str, object]] = []

    for sample, spec in SERIES.items():
        grain_modulus = float(sheet[spec["grain_modulus_cell"]].value)
        for row in range(32, 48):
            pressure = sheet.cell(row, spec["pressure_col"]).value
            drained_modulus = sheet.cell(row, spec["modulus_col"]).value
            alpha = sheet.cell(row, spec["alpha_col"]).value
            # The SP30 workbook contains one modulus/alpha cell without an x value.
            # Excel does not plot it in Figure 3c, so it is excluded here.
            if pressure is None or drained_modulus is None or alpha is None:
                continue
            recomputed = 1.0 - float(drained_modulus) / grain_modulus
            records.append(
                {
                    "sample": sample,
                    "treatment": spec["label"],
                    "pressure_mpa": float(pressure),
                    "drained_bulk_modulus_gpa": float(drained_modulus),
                    "unjacketed_bulk_modulus_gpa": grain_modulus,
                    "biot_coefficient_reported": float(alpha),
                    "biot_coefficient_recomputed": recomputed,
                    "recompute_error": float(alpha) - recomputed,
                    "source_sheet": "Integrated",
                    "source_row": row,
                }
            )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)

    counts = {sample: sum(row["sample"] == sample for row in records) for sample in SERIES}
    print(f"wrote {len(records)} Figure 3c points to {args.output}: {counts}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

